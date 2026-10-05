# -*- coding: utf-8 -*-
"""
Build the human sentiment-annotation workbooks for the 270-item validation sample.

Same items as the stance validation, so stance and tone are judged on identical
comments. Two annotators each label all 270 independently, so agreement can be
reported per platform (135 items each); a third resolves only the items they
disagree on. Each workbook is shuffled with its own fixed seed, so the order
differs between annotators and from the stance workbooks.

Round 2: round 1 showed only the subreddit/post title or a bare video ID, and
annotators found many comments - most of them replies - unreadable without the
comment they answer. Round 2 adds the video title and the parent comment
(build_validation_context.py) and a new item order. Round-1 sheets are kept in
validation/sentiment_round1/ and are not used.

Run:  python 06_revision/make_sentiment_workbooks.py
Out:  06_revision/validation/sentiment_<name>.xlsx
"""
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from make_annotation_workbooks import (BLIND, BOX, FIRST, FONT, HDR_FILL, HDR_ROW,
                                       INK, INK_B, INPUT_FILL, MUTED, TITLE,
                                       TODO_FILL, VAL)

HERE = Path(__file__).resolve().parent
SEED = 20261106          # round 2: new order, not the round-1 order
ANNOTATORS = ("Akshat", "Raghav")       # adjudication of disagreements: Arsh

LABELS = ("POS", "NEG", "NEU")

COLS = [
    ("item_id", 10),
    ("platform", 11),
    ("context", 30),
    ("replying to", 46),
    ("comment text", 66),
    ("sentiment", 13),
    ("notes", 28),
]

LABEL_DEFS = [
    ("POS", "Positive",
     "The overall tone is favourable: approval, praise, hope, gratitude, relief, "
     "support, warmth, humour that is friendly.",
     "“Glad the hostages are home, this made my day.”"),
    ("NEG", "Negative",
     "The overall tone is unfavourable: anger, contempt, grief, fear, disgust, "
     "mockery, blame, despair, insults.",
     "“These people are monsters and everyone defending them should be ashamed.”"),
    ("NEU", "Neutral",
     "No clear emotional tone: factual statements, plain questions, requests, "
     "neutral reporting - or positive and negative so evenly balanced that "
     "neither dominates.",
     "“Does anyone have the source for the casualty numbers?”"),
]

RULES = [
    "Context (column C) and the comment being replied to (column D) are there to "
    "help you READ the comment - what it refers to, whether it is sarcastic. Label "
    "the tone of the comment in column E only, never the tone of the post, video or "
    "parent comment. A calm reply to an angry comment is NEU or POS, not NEG.",
    "Label the TONE of the wording, not the side it takes. A pro-Israel or a "
    "pro-Palestine comment can each be positive, negative or neutral. Do not "
    "use your stance judgement from the earlier sheet.",
    "Do not label whether you agree with it, or whether the event it describes is "
    "good or bad. “Ceasefire talks resumed today.” is NEU even though a ceasefire is "
    "welcome news.",
    "Anger at a third party is still NEG. “Hamas are butchers” and “the IDF are "
    "butchers” are both NEG, whatever the author's side.",
    "Sarcasm: label the tone actually meant. “Great job, another hospital bombed” is "
    "NEG.",
    "Mixed comments: pick the tone that dominates. If it is genuinely balanced, NEU.",
    "Off-topic comments still get a tone label. Every row needs one.",
    "Very short or fragmentary comments: label what is there. If there is no tone to "
    "read, NEU.",
    "Work independently and do not discuss items with the other annotator until both "
    "of you are done. Agreement between you is what is being measured.",
    "Go with your first considered reading. Aim for about 10 seconds per comment and "
    "do not agonise; leave a note on anything genuinely hard.",
]


def guidelines_sheet(wb, who, n):
    ws = wb.create_sheet("Guidelines", 0)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", (6, 22, 74, 52)):
        ws.column_dimensions[col].width = w

    ws["A1"] = "Sentiment annotation - guidelines"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Annotator {who}  |  {n} items  |  "
                "Israel-Hamas discourse validation sample")
    ws["A2"].font = MUTED

    ws["A4"] = "What to do"
    ws["A4"].font = INK_B
    steps = [
        "Go to the Annotate sheet.",
        "For each row, read the context (C) and the comment being replied to (D), "
        "then the comment itself (E), and choose a label in column F (yellow). "
        "Column F has a dropdown - POS, NEG or NEU.",
        "Optionally add a note in column G.",
        "Rows you have not labelled stay highlighted. The counter at the top of the "
        "sheet tracks progress.",
        f"When all {n} are done, save the file (keep the .xlsx name) and hand it back.",
    ]
    r = 5
    for i, s in enumerate(steps, 1):
        ws.cell(r, 1, f"{i}.").font = INK
        c = ws.cell(r, 2, s)
        c.font = INK
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    ws.cell(r, 1, "Labels").font = INK_B
    r += 1
    for j, h in enumerate(("", "Meaning", "When it applies", "Example"), 1):
        c = ws.cell(r, j, h)
        c.font = INK_B
        c.fill = HDR_FILL
        c.border = BOX
    r += 1
    for code, meaning, desc, example in LABEL_DEFS:
        for j, v in enumerate((code, meaning, desc, example), 1):
            c = ws.cell(r, j, v)
            c.font = INK_B if j == 1 else INK
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOX
        ws.row_dimensions[r].height = 58
        r += 1

    r += 1
    ws.cell(r, 1, "Rules").font = INK_B
    r += 1
    for s in RULES:
        ws.cell(r, 1, "•").font = INK
        c = ws.cell(r, 2, s)
        c.font = INK
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.cell(r, 1, "Worked examples").font = INK_B
    r += 1
    examples = [
        ("“Stay strong Gaza, the whole world is with you ❤️”", "POS",
         "supportive and warm - the side taken does not matter"),
        ("“Proud of our soldiers, bring them all home safe”", "POS",
         "pride and hope"),
        ("“Absolutely disgusting. How is this allowed in 2024?”", "NEG",
         "outrage"),
        ("“lol imagine believing the ministry of health numbers”", "NEG",
         "mockery"),
        ("“The UN vote is scheduled for Friday.”", "NEU", "plain fact"),
        ("“Which channel is this from?”", "NEU", "plain question"),
    ]
    for text, lab, why in examples:
        ws.cell(r, 2, text).font = INK
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        c = ws.cell(r, 3, f"{lab}  -  {why}")
        c.font = INK_B
        c.fill = INPUT_FILL
        ws.row_dimensions[r].height = 18
        r += 1
    ws.cell(r + 1, 2, "Only the yellow cells (sentiment, and optionally notes) "
                      "are yours to fill in.").font = MUTED
    return ws


def annotate_sheet(wb, who, df):
    ws = wb.create_sheet("Annotate", 1)
    n = len(df)
    last = FIRST + n - 1

    for i, (name, width) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws["A1"] = f"Sentiment - annotator {who}"
    ws["A1"].font = TITLE
    ws["A2"] = ("Choose POS / NEG / NEU in the yellow sentiment column. "
                "Unlabelled rows stay highlighted.")
    ws["A2"].font = MUTED

    lab_col = get_column_letter(6)
    rng = f"${lab_col}${FIRST}:${lab_col}${last}"
    ws["F1"] = "labelled"
    ws["F1"].font = MUTED
    ws["G1"] = "=" + "+".join(f'COUNTIF({rng},"{l}")' for l in LABELS)
    ws["G1"].font = Font(name=FONT, size=13, bold=True)
    ws["F2"] = "remaining"
    ws["F2"].font = MUTED
    ws["G2"] = f"=COUNTA($A${FIRST}:$A${last})-$G$1"
    ws["G2"].font = INK_B

    for i, (name, _) in enumerate(COLS, 1):
        c = ws.cell(HDR_ROW, i, name)
        c.font = INK_B
        c.fill = HDR_FILL
        c.border = BOX
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[HDR_ROW].height = 20

    top_wrap = Alignment(wrap_text=True, vertical="top")
    for k, row in enumerate(df.itertuples(index=False), start=FIRST):
        ctx = row.context_subreddit or ""
        if ctx:
            ctx = f"r/{ctx}"
            if row.context_post_title:
                ctx += f"\n“{row.context_post_title}”"
        elif row.context_video_id:
            ctx = "YouTube video"
            if row.video_title:
                ctx += f"\n“{row.video_title}”"
        parent = str(row.parent_text) if row.is_reply else ""
        if row.is_reply and not parent:
            parent = "(reply - parent comment unavailable)"
        if parent and parent[0] in "=+@-":
            parent = "'" + parent

        text = str(row.text)
        if text[:1] in "=+@-":         # never let a comment parse as a formula
            text = "'" + text

        values = [row.item_id, row.platform, ctx, parent, text, None, None]
        for i, v in enumerate(values, 1):
            c = ws.cell(k, i, v)
            c.font = INK
            c.alignment = top_wrap
            c.border = BOX
            if i in (6, 7):
                c.fill = INPUT_FILL
        ws.cell(k, 6).alignment = Alignment(horizontal="center", vertical="center")
        lines = max(math.ceil(len(text) / 80), math.ceil(len(parent) / 55),
                    ctx.count("\n") + 2)
        ws.row_dimensions[k].height = 15 * min(max(lines, 2), 9)

    dv = DataValidation(
        type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True,
        showDropDown=False, errorTitle="Not a valid label",
        error="Use POS (positive), NEG (negative) or NEU (neutral).",
        promptTitle="Sentiment", prompt="POS / NEG / NEU")
    dv.error_style = "stop"
    ws.add_data_validation(dv)
    dv.add(f"{lab_col}{FIRST}:{lab_col}{last}")

    ws.conditional_formatting.add(
        f"A{FIRST}:G{last}",
        FormulaRule(formula=[f"AND($A{FIRST}<>\"\",$F{FIRST}=\"\")"],
                    fill=TODO_FILL, stopIfTrue=False))

    ws.freeze_panes = f"A{FIRST}"
    ws.auto_filter.ref = f"A{HDR_ROW}:G{last}"
    return ws


def write(who, df):
    wb = Workbook()
    wb.remove(wb.active)
    guidelines_sheet(wb, who, len(df))
    annotate_sheet(wb, who, df)
    wb.active = 1
    out = VAL / f"sentiment_{who}.xlsx"
    wb.save(out)
    print(f"  wrote {out.name}  ({len(df)} items, {out.stat().st_size/1024:.0f} KB)")


def main():
    df = pd.read_csv(BLIND).fillna("")
    ctx = pd.read_parquet(HERE / "outputs" / "validation_context.parquet")
    df = df.merge(ctx, on="item_id", how="left")
    assert df["is_reply"].notna().all()
    print(f"{len(df)} items from {BLIND.name}")
    for i, who in enumerate(ANNOTATORS):
        write(who, df.sample(frac=1, random_state=SEED + i).reset_index(drop=True))


if __name__ == "__main__":
    main()
